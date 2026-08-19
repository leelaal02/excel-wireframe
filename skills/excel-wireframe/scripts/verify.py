# -*- coding: utf-8 -*-
"""생성된 pptx를 다시 파싱해 기대와 대조한다.

만들었다는 사실만으로는 부족하다. 도형 이름이 안 맞거나 rId가 깨지면
파일은 생기지만 내용이 비어 있을 수 있다.
"""
from __future__ import annotations

import re
from pathlib import Path

from common import resolve_template_path
from pptx import Presentation
from pptx_scan import scan_presentation
from slide_fill import collect_tables
from slide_layout import meta_slot_name


def label_for(meta_cfg: dict | None, key: str) -> str | None:
    """메타 표에서 이 값을 담는 라벨. `{"화면명": "title"}`을 거꾸로 찾는다."""
    for label, value_key in ((meta_cfg or {}).get("labels") or {}).items():
        if value_key == key:
            return label
    return None


def _meta_slot_values(slide: dict, meta_cfg: dict | None, key: str) -> list[str]:
    """메타 표 위에 얹힌 글자 자리에서 값을 모은다.

    기본 템플릿은 화면명·화면ID를 도형이 아니라 상단 메타 표의 칸 자리에 넣는다.
    그 자리를 안 보면 값이 들어갔는지 확인할 방법이 없다.
    """
    label = label_for(meta_cfg, key)
    if not label:
        return []
    want = meta_slot_name(label)
    return [sh["text"] for sh in slide["shapes"] if sh["name"] == want]


def _title_texts(slide: dict, title_name: str | None,
                 meta_cfg: dict | None = None) -> list[str]:
    """슬라이드에서 화면명 판단에 쓸 텍스트를 뽑는다.

    제목 도형 이름이 지정되면 그 도형을 보고, 메타 표가 화면명을 담으면 그
    칸도 본다. 둘 다 없으면(매핑이 제목을 지정하지 않으면) 어느 도형이든
    검사하던 예전 방식으로 물러선다 — 크래시를 피하기 위함이지, 정확도를
    위한 것은 아니다.
    """
    texts = []
    if title_name:
        texts += [sh["text"] for sh in slide["shapes"] if sh["name"] == title_name]
    texts += _meta_slot_values(slide, meta_cfg, "title")
    if title_name or texts:
        return texts
    return [sh["text"] for sh in slide["shapes"] if sh["text"]]


def _matches_screen_name(text: str, screen_name: str) -> bool:
    """제목 텍스트가 정확히 이 화면 것인지 판단한다.

    예전엔 부분 문자열 포함 여부만 봤다 — "목록"과 "이용기관 목록"처럼 한
    화면명이 다른 화면명의 부분 문자열이면 한 슬라이드가 두 화면 모두에
    매치됐다. 화면명 반영/이미지 배치 검사는 "매치가 있는가"만 봐서 이
    과다매치가 무해했지만, 상세 항목 수 검사는 매치된 슬라이드를 전부
    합산하므로 과다매치가 그대로 오탐이 된다. page_title()이 만드는 꼴은
    "이름" 또는 "이름 (i/n)" 뿐이므로 끝을 고정하면 정확히 잡아낼 수 있다.
    """
    pattern = r"^%s(\s*\(\d+/\d+\))?$" % re.escape(screen_name)
    return re.match(pattern, text.strip()) is not None


def _slides_for_screen(slides: list[dict], title_name: str | None,
                       screen_name: str, meta_cfg: dict | None = None) -> list[dict]:
    """화면명이 제목(또는 대체 텍스트)과 정확히 일치하는 슬라이드를 모두 찾는다.

    분할된 화면은 '이름 (1/2)', '이름 (2/2)'처럼 슬라이드가 여러 장이므로
    첫 번째 매치만 보면 안 된다.
    """
    return [
        s for s in slides
        if any(_matches_screen_name(t, screen_name)
               for t in _title_texts(s, title_name, meta_cfg))
    ]


def _slides_by_screen_id(slides: list[dict], id_name: str | None,
                         screen_id: str, meta_cfg: dict | None = None) -> list[dict]:
    """화면ID 도형에 이 화면의 ID가 찍힌 슬라이드를 모두 찾는다.

    분할된 화면은 장마다 같은 ID를 달고 나오므로 여러 장이 잡히는 게 정상이다.
    """
    return [
        s for s in slides
        if any(sh["name"] == id_name and sh["text"].strip() == screen_id
               for sh in s["shapes"])
        or screen_id in [v.strip() for v in _meta_slot_values(s, meta_cfg, "screen_id")]
    ]


def verify_output(out_path: Path, screens_data: dict, mapping: dict,
                  expected_slides: int, work_dir: Path | None = None) -> dict:
    report = scan_presentation(Path(out_path))
    slides = report["slides"]
    checks: list[dict] = []

    checks.append(
        {
            "name": "슬라이드 수",
            "ok": len(slides) == expected_slides,
            "detail": "기대 %d장, 실제 %d장" % (expected_slides, len(slides)),
        }
    )

    title_name = mapping.get("template", {}).get("shapes", {}).get("title")
    id_name = mapping.get("template", {}).get("shapes", {}).get("screen_id")
    meta_cfg = mapping.get("template", {}).get("meta_table")
    screens = screens_data.get("screens", [])

    def slides_of(scr: dict) -> list[dict]:
        """화면의 **내용**(이미지·상세)을 검사할 슬라이드를 고른다.

        화면명은 실무에서 겹친다 — '목록'·'상세'·'등록'은 메뉴마다 되풀이되고
        구분은 화면ID가 한다. 이름으로 고르면 동명 화면끼리 서로의 슬라이드를
        집어삼켜 상세 건수가 합산되므로, 화면ID 도형이 지정돼 있으면 그 값으로
        짚는다. build가 그 도형에 넣는 것과 같은 값이라 정확히 맞는다.

        ID로 한 장도 못 찾으면 이름 매칭으로 물러선다 — 템플릿에 화면ID 도형이
        없으면(shape-not-found) ID가 애초에 안 찍히므로, 그 경우까지 "슬라이드를
        찾지 못함"으로 몰면 멀쩡한 결과물이 실패한다.

        '화면명 반영' 검사는 이 함수를 쓰지 않는다. 그쪽은 제목에 화면명이
        들어갔는지를 보는 검사라 이름으로 찾는 것이 곧 검사 내용이다.
        """
        if (id_name or meta_cfg) and scr.get("id"):
            found = _slides_by_screen_id(slides, id_name, scr["id"], meta_cfg)
            if found:
                return found
        return _slides_for_screen(slides, title_name, scr["name"], meta_cfg)

    # expected_slides는 screens_data 자체에서 파생되므로, 화면이 통째로 0개면
    # 슬라이드 수 검사도 "기대 0장, 실제 0장"으로 공허하게 통과하고 나머지
    # 검사도 빈 목록을 돌며 전부 통과한다. sheet_include 오타처럼 매핑이 완전히
    # 잘못된 가장 흔한 사고가 "검증: 통과"로 보고되는 사고를 막는다.
    checks.append(
        {
            "name": "화면 데이터",
            "ok": bool(screens),
            "detail": "화면 %d개" % len(screens) if screens
            else "화면이 0개입니다 — mapping.json의 excel.sheet_include"
                 "(sheet-per-screen) 또는 excel 레이아웃 설정이 실제 Excel과 "
                 "맞지 않을 가능성이 큽니다",
        }
    )

    missing_titles = [
        scr["name"] for scr in screens
        if not _slides_for_screen(slides, title_name, scr["name"], meta_cfg)
    ]
    checks.append(
        {
            "name": "화면명 반영",
            "ok": not missing_titles,
            "detail": "누락 없음" if not missing_titles
            else "누락: %s" % ", ".join(missing_titles),
        }
    )

    image_issues: list[str] = []
    for scr in screens:
        matches = slides_of(scr)
        if not matches:
            image_issues.append("%s: 화면에 해당하는 슬라이드를 찾지 못함" % scr["id"])
            continue
        if scr.get("images"):
            has_pic = any(
                any("PICTURE" in sh["shape_type"] for sh in s["shapes"])
                for s in matches
            )
            if not has_pic:
                image_issues.append("%s: 이미지가 배치되지 않음" % scr["id"])
    checks.append(
        {
            "name": "이미지 배치",
            "ok": not image_issues,
            "detail": "문제 없음" if not image_issues else "; ".join(image_issues),
        }
    )

    # 설계 스펙의 5개 검증 항목 중 "상세 항목 수와 표에 채워진 슬롯 수가
    # 일치하는가"에 해당한다. detail_tables 이름이 틀리면 collect_tables가
    # 빈 목록을 반환해 fill_slots가 아예 안 돌고, 표는 템플릿의 예시 텍스트
    # 그대로 남는다 — 앞의 네 검사는 이 상황에서도 전부 통과하므로 이 검사가
    # 유일하게 잡아낸다.
    opts = mapping.get("options", {})
    clear_unused = bool(opts.get("clear_unused_slots", True))
    if not clear_unused:
        # clear_unused_slots가 꺼져 있으면 안 쓰는 슬롯은 템플릿의 예시
        # 텍스트를 일부러 그대로 남긴다 — "채워진 슬롯 수"가 애초에 의미가
        # 없으므로 검사를 걸러 오탐을 만들지 않는다.
        checks.append(
            {
                "name": "상세 항목 수",
                "ok": True,
                "detail": "options.clear_unused_slots가 꺼져 있어 검사를 건너뜁니다"
                          "(안 쓰는 슬롯이 예시 텍스트를 그대로 유지하므로 채워진"
                          " 슬롯 수가 상세 건수를 나타내지 않습니다)",
            }
        )
    else:
        tpl_cfg = mapping.get("template", {})
        table_cols = tpl_cfg.get("table_columns", {"no": 0, "text": 1})
        no_col = int(table_cols.get("no", 0))
        text_col = int(table_cols.get("text", 1))
        detail_names = tpl_cfg.get("shapes", {}).get("detail_tables")
        prs = Presentation(str(out_path))

        detail_issues: list[str] = []
        for scr in screens:
            matches = slides_of(scr)
            if not matches:
                continue  # 화면명 반영 검사가 이미 이 화면을 짚었다
            expected = len(scr.get("details", []))
            actual = 0
            for s in matches:
                slide = prs.slides[s["index"]]
                for t in collect_tables(slide, detail_names):
                    table = t.table
                    n_cols = len(table.columns)
                    for r in range(len(table.rows)):
                        # 번호 칸 또는 설명 칸 중 하나라도 차 있으면 슬롯이
                        # 채워진 것으로 본다. desc가 빈 문자열인 상세 행도
                        # _read_details는 no/요소명 등 다른 칸이 차 있으면
                        # 그대로 통과시키므로, 설명 칸만 보면 그런 상세가
                        # "슬롯 미달"로 오탐된다 — 번호는 Excel에서 왔으니
                        # 설명이 비어도 항상 있다.
                        no_filled = no_col < n_cols and table.cell(r, no_col).text.strip()
                        text_filled = text_col < n_cols and table.cell(r, text_col).text.strip()
                        if no_filled or text_filled:
                            actual += 1
            if actual != expected:
                detail_issues.append(
                    "%s: 상세 %d건, 표에 채워진 슬롯 %d개" % (scr["id"], expected, actual)
                )
        checks.append(
            {
                "name": "상세 항목 수",
                "ok": not detail_issues,
                "detail": "문제 없음" if not detail_issues else "; ".join(detail_issues),
            }
        )

    template = resolve_template_path(mapping["template"], work_dir)
    if template.exists():
        tpl_report = scan_presentation(template)
        same = (
            tpl_report["slide_width"] == report["slide_width"]
            and tpl_report["slide_height"] == report["slide_height"]
        )
        checks.append(
            {
                "name": "슬라이드 크기",
                "ok": same,
                "detail": "%.2f x %.2f in" % tuple(report["slide_size_in"]),
            }
        )
    else:
        checks.append(
            {
                "name": "슬라이드 크기",
                "ok": False,
                "detail": "템플릿 파일을 찾을 수 없어 비교 불가: %s" % template,
            }
        )

    return {"ok": all(c["ok"] for c in checks), "checks": checks}
