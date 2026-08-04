# -*- coding: utf-8 -*-
"""슬라이드 복제. python-pptx에 복제 API가 없어 직접 구현한다.

관계(rels)를 새 슬라이드에 재등록할 때 rId를 지정할 수 없고 자동 할당되므로,
복제한 XML 안의 r:embed / r:id 등을 새 rId로 치환해야 한다. 이 재매핑을 빼면
그림이 사라지거나 파일이 손상된다.
"""
from __future__ import annotations

import copy

from pptx.opc.constants import RELATIONSHIP_TYPE as RT

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
R_ATTRS = [
    "{%s}%s" % (R_NS, n)
    for n in ("id", "embed", "link", "pict", "dm", "lo", "qs", "cs")
]
SKIP = {RT.SLIDE_LAYOUT, RT.SLIDE_MASTER}


def clone_slide(prs, src):
    """src 슬라이드를 프레젠테이션 끝에 복제하고 새 슬라이드를 반환한다."""
    new = prs.slides.add_slide(src.slide_layout)

    # 레이아웃이 자동 삽입한 플레이스홀더를 걷어낸다. 원본 도형만 남겨야
    # 템플릿 모양이 정확히 재현된다.
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)

    rid_map: dict[str, str] = {}
    for old_rid, rel in src.part.rels.items():
        if rel.reltype in SKIP:
            continue
        if rel.is_external:
            new_rid = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = new.part.rels.get_or_add(rel.reltype, rel.target_part)
        rid_map[old_rid] = new_rid

    spTree = new.shapes._spTree
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in R_ATTRS:
                v = node.get(attr)
                if v is not None and v in rid_map:
                    node.set(attr, rid_map[v])
        spTree.append(el)

    return new
